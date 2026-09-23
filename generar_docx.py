# -*- coding: utf-8 -*-
"""Convierte el .md del informe a .docx (python-docx), emulando el estilo del PDF."""
import re
import pathlib
from urllib.parse import unquote

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "Arquitectura para sistemas IA.md"
OUT = "EP1_Equipo1_HigueraDuranSalazar.docx"

TOK = re.compile(r"(`[^`\n]+`)|(\*\*[^*\n]+\*\*)|(\*[^*\n]+\*|_[^_\n]+_)")

doc = Document()

# ---- Página A4 y márgenes (espejo del PDF) ----
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.1)
sec.left_margin = sec.right_margin = Cm(1.55)

# ---- Estilos base ----
normal = doc.styles["Normal"]
normal.font.name = "Georgia"
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.23
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

HEADING_STYLES = {"#": ("Title", 18), "##": ("Heading 1", 13),
                  "###": ("Heading 2", 11.5), "####": ("Heading 3", 10.5)}
for _, (name, size) in HEADING_STYLES.items():
    st = doc.styles[name]
    st.font.name = "Georgia"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    st.paragraph_format.space_before = Pt(8)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.keep_with_next = True


def _run(p, text, size=None, bold=None, italic=None):
    r = p.add_run(text)
    if size:
        r.font.size = size
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    return r


def add_inline(p, text, size=None, base_bold=None, base_italic=None):
    """Escribe texto con **negrita**, *cursiva*, _cursiva_ y `codigo` (anidables)."""
    pos = 0
    for m in TOK.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], size, bold=base_bold, italic=base_italic)
        if m.group(1):                      # `codigo`
            r = _run(p, m.group(1)[1:-1], size, bold=base_bold, italic=base_italic)
            r.font.name = "Consolas"
        elif m.group(2):                    # **negrita** (puede anidar `codigo`)
            add_inline(p, m.group(2)[2:-2], size, base_bold=True, base_italic=base_italic)
        else:                               # *cursiva* / _cursiva_
            add_inline(p, m.group(3)[1:-1], size, base_bold=base_bold, base_italic=True)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], size, bold=base_bold, italic=base_italic)


def add_table(block):
    rows = []
    for bl in block:
        if re.match(r"^\|[\s|:-]+\|\s*$", bl):      # separador |---|
            continue
        rows.append([c.strip() for c in bl.strip().strip("|").split("|")])
    if not rows:
        return
    header, body = rows[0], rows[1:]
    ncols = len(header)
    tbl = doc.add_table(rows=1 + len(body), cols=ncols)
    tbl.style = "Table Grid"
    for j, h in enumerate(header):                   # cabecera en negrita
        p = tbl.cell(0, j).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, h, size=Pt(9))
        for r in p.runs:
            r.bold = True
    for ri, row in enumerate(body, start=1):         # filas
        row = row + [""] * (ncols - len(row))
        for j in range(ncols):
            p = tbl.cell(ri, j).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, row[j], size=Pt(9))


lines = pathlib.Path(SRC).read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    i += 1

    if not line.strip():
        continue
    if "page-break-after" in line:          # saltos de portada/índice
        doc.add_page_break()
        continue
    if re.match(r"^---+$", line):           # separadores horizontales
        continue

    if line.lstrip().startswith("|"):       # tablas (bloque completo)
        block = [line]
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            block.append(lines[i])
            i += 1
        add_table(block)
        continue

    m = re.match(r"^(#{1,4})\s+(.*)$", line)
    if m:
        style, _ = HEADING_STYLES[m.group(1)]
        p = doc.add_paragraph(style=style)
        add_inline(p, m.group(2).strip())
        continue

    m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", line)
    if m:                                    # imagen
        doc.add_picture(unquote(m.group(2)), width=Cm(14.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue

    m = re.match(r"^(\s*)-\s+(.*)$", line)   # viñetas
    if m:
        level = min(len(m.group(1)), 4) // 2
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6 + 0.7 * level)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, m.group(2))
        continue

    m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)  # pasos numerados (literal)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(p, f"{m.group(2)}. {m.group(3)}")
        continue

    p = doc.add_paragraph()                 # párrafo normal
    add_inline(p, line)

doc.save(OUT)

# ---- Verificación tras guardar ----
chk = Document(OUT)
n_img = len(chk.inline_shapes)
n_tbl = len(chk.tables)
cells = sum(len(t.rows) * len(t.columns) for t in chk.tables)
print(f"DOCX guardado: {OUT}")
print(f"parrafos: {len(chk.paragraphs)} | tablas: {n_tbl} ({cells} celdas) | imagenes: {n_img}")
